"""Filesystem regressions; crypto is simulated, not encryption validation."""

import asyncio
import builtins
import errno
import hashlib
import json
import os
from pathlib import Path
import shutil
import stat
import sys
from types import SimpleNamespace
from typing import BinaryIO
import zipfile

from docx import Document
import pytest

from word_document_server.core import footnotes, protection, unprotect
from word_document_server.tools import footnote_tools
from word_document_server.utils.file_utils import create_document_copy


def _deny_metadata(monkeypatch: pytest.MonkeyPatch) -> None:
    def denied(*args: object, **kwargs: object) -> None:
        raise PermissionError(errno.EPERM, "metadata changes denied")

    for name in ("chmod", "chown", "utime"):
        if hasattr(os, name):
            monkeypatch.setattr(os, name, denied)
    monkeypatch.setattr(shutil, "copystat", denied)


def test_copy_bytes_without_metadata(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "source.docx"
    payload = bytes(range(256)) * 1024
    source.write_bytes(payload)
    _deny_metadata(monkeypatch)

    success, message, copied = create_document_copy(str(source))

    assert success, message
    assert copied == str(tmp_path / "source_copy.docx")
    assert Path(copied).read_bytes() == payload
    assert source.read_bytes() == payload


@pytest.mark.skipif(os.name != "posix", reason="POSIX permission bits")
@pytest.mark.parametrize("source_mode", [0o600, 0o640])
def test_copy_is_private_from_creation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, source_mode: int
) -> None:
    source = tmp_path / "private.docx"
    destination = tmp_path / "copy.docx"
    source.write_bytes(b"private document")
    source.chmod(source_mode)
    _deny_metadata(monkeypatch)
    observed_modes = []

    real_open = builtins.open

    def observe_open(path: str, mode: str = "r", **kwargs: object) -> BinaryIO:
        stream = real_open(path, mode, **kwargs)
        if path == str(destination):
            observed_modes.append(stat.S_IMODE(os.fstat(stream.fileno()).st_mode))
        return stream

    monkeypatch.setattr(builtins, "open", observe_open)
    previous_umask = os.umask(0o022)
    try:
        success, message, _ = create_document_copy(str(source), str(destination))
    finally:
        os.umask(previous_umask)

    assert success, message
    assert destination.read_bytes() == source.read_bytes()
    assert stat.S_IMODE(destination.stat().st_mode) & 0o077 == 0
    assert observed_modes
    assert all(mode & 0o077 == 0 for mode in observed_modes)


@pytest.mark.parametrize("alias", ["same_path", "hard_link"])
def test_copy_rejects_source_alias_without_truncating(
    tmp_path: Path, alias: str
) -> None:
    source = tmp_path / "source.docx"
    source.write_bytes(b"keep original")
    destination = source
    if alias == "hard_link":
        destination = tmp_path / "alias.docx"
        os.link(source, destination)

    success, _, _ = create_document_copy(str(source), str(destination))

    assert not success
    assert source.read_bytes() == b"keep original"


@pytest.mark.skipif(os.name != "posix", reason="POSIX permission bits")
def test_overwrite_retains_destination_permissions_without_metadata_calls(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "source.docx"
    destination = tmp_path / "existing.docx"
    source.write_bytes(b"replacement")
    destination.write_bytes(b"old longer content")
    destination.chmod(0o660)
    _deny_metadata(monkeypatch)

    success, message, _ = create_document_copy(str(source), str(destination))

    assert success, message
    assert destination.read_bytes() == b"replacement"
    assert stat.S_IMODE(destination.stat().st_mode) == 0o660


@pytest.mark.skipif(os.name != "posix", reason="POSIX permission bits")
@pytest.mark.parametrize("operation", ["core_add", "core_delete", "after", "before", "delete"])
@pytest.mark.parametrize("valid_target", [False, True])
def test_footnote_output_copy_is_private(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, operation: str, valid_target: bool
) -> None:
    source = tmp_path / "source.docx"
    destination = tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("target paragraph")
    doc.save(source)
    if "delete" in operation:
        success, message, _ = footnotes.add_footnote_robust(
            str(source), search_text="target", footnote_text="note"
        )
        assert success, message
    source.chmod(0o600)
    original = source.read_bytes()
    _deny_metadata(monkeypatch)
    search_text = "target" if valid_target else "missing text"
    previous_umask = os.umask(0o022)
    try:
        if operation == "core_add":
            success, message, _ = footnotes.add_footnote_robust(
                str(source), search_text=search_text, footnote_text="note",
                output_filename=str(destination),
            )
        elif operation == "core_delete":
            success, message, _ = footnotes.delete_footnote_robust(
                str(source), search_text=search_text, output_filename=str(destination)
            )
        else:
            function = {
                "after": footnote_tools.add_footnote_after_text_robust,
                "before": footnote_tools.add_footnote_before_text_robust,
                "delete": footnote_tools.delete_footnote_from_document_robust,
            }[operation]
            kwargs = {} if operation == "delete" else {"footnote_text": "note"}
            message = asyncio.run(function(
                str(source), search_text=search_text,
                output_filename=str(destination), **kwargs,
            ))
            success = "Successfully" in message
    finally:
        os.umask(previous_umask)

    assert success == valid_target, message
    assert source.read_bytes() == original
    assert stat.S_IMODE(destination.stat().st_mode) & 0o077 == 0
    if not valid_target:
        assert destination.read_bytes() == original
    else:
        with zipfile.ZipFile(destination) as archive:
            document_xml = archive.read("word/document.xml")
            assert (b"footnoteReference" in document_xml) == ("delete" not in operation)


@pytest.mark.parametrize("operation", ["add", "delete"])
def test_footnote_replace_failure_keeps_original_and_cleans_staging(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, operation: str
) -> None:
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("target paragraph")
    doc.save(source)
    success, message, _ = footnotes.add_footnote_robust(
        str(source), search_text="target", footnote_text="note"
    )
    assert success, message
    original = source.read_bytes()
    before = source.stat()
    staged_paths = []

    def denied(staged: str, destination: str) -> None:
        staged_paths.append(Path(staged))
        raise PermissionError(errno.EPERM, "replacement denied")

    monkeypatch.setattr(os, "replace", denied)
    _deny_metadata(monkeypatch)
    function = footnotes.add_footnote_robust if operation == "add" else footnotes.delete_footnote_robust
    success, message, _ = function(str(source), search_text="target")

    assert not success
    assert "replacement denied" in message
    assert len(staged_paths) == 1
    assert staged_paths[0].parent == source.parent
    assert set(tmp_path.iterdir()) == {source}
    assert source.read_bytes() == original
    after = source.stat()
    assert (after.st_ino, after.st_mode, after.st_mtime_ns) == (
        before.st_ino, before.st_mode, before.st_mtime_ns
    )


@pytest.mark.parametrize("operation", ["protect", "unprotect"])
@pytest.mark.parametrize("replace_fails", [False, True])
def test_protection_staging_and_replacement(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, operation: str, replace_fails: bool
) -> None:
    source = tmp_path / "document.docx"
    original = b"original document bytes"
    transformed = b"simulated crypto output"
    source.write_bytes(original)
    os.utime(source, ns=(1_600_000_000_000_000_000,) * 2)
    before = source.stat()
    metadata = source.with_suffix(".protection")
    password = "test-only-password"
    password_hash = hashlib.sha256(password.encode()).hexdigest()
    metadata_bytes = json.dumps({
        "true_encryption": True, "password_hash": password_hash,
    }).encode()
    if operation == "unprotect":
        metadata.write_bytes(metadata_bytes)
    input_handles = []
    output_handles = []

    class SimulatedOfficeFile:
        """Deliberately accepts the legacy API; no claim of real encryption."""

        def __init__(self, stream: BinaryIO) -> None:
            assert stream.read() == original
            input_handles.append(stream)

        def load_key(self, password: str) -> None:
            assert password == "test-only-password"

        def encrypt(self, stream: BinaryIO) -> None:
            output_handles.append(stream)
            stream.write(transformed)

        decrypt = encrypt

    monkeypatch.setitem(sys.modules, "msoffcrypto", SimpleNamespace(OfficeFile=SimulatedOfficeFile))
    real_replace = os.replace
    staged_paths = []
    staged_bytes = []
    closed_at_replace = []

    def replace(staged: str, destination: str) -> None:
        staged_path = Path(staged)
        staged_paths.append(staged_path)
        staged_bytes.append(staged_path.read_bytes())
        closed_at_replace.append(all(handle.closed for handle in input_handles + output_handles))
        if any(not handle.closed for handle in input_handles):
            raise PermissionError(errno.EACCES, "simulated Windows sharing violation")
        if replace_fails:
            raise PermissionError(errno.EPERM, "replacement denied")
        real_replace(staged, destination)

    monkeypatch.setattr(os, "replace", replace)
    _deny_metadata(monkeypatch)
    if operation == "protect":
        success = protection.add_protection_info(
            str(source), "password", password_hash, raw_password=password
        )
    else:
        success, _ = unprotect.remove_protection_info(str(source), password)

    assert success == (not replace_fails)
    assert len(staged_paths) == 1
    assert staged_paths[0].parent == source.parent
    assert staged_bytes == [transformed]
    assert closed_at_replace == [True]
    assert not staged_paths[0].exists()
    assert all(handle.closed for handle in input_handles + output_handles)
    assert source.read_bytes() == (original if replace_fails else transformed)
    if replace_fails:
        after = source.stat()
        assert (after.st_ino, after.st_mode, after.st_uid, after.st_gid, after.st_mtime_ns) == (
            before.st_ino, before.st_mode, before.st_uid, before.st_gid, before.st_mtime_ns
        )
        if operation == "unprotect":
            assert metadata.read_bytes() == metadata_bytes
        else:
            assert not json.loads(metadata.read_text()).get("true_encryption", False)
    elif operation == "unprotect":
        assert not metadata.exists()
    else:
        assert json.loads(metadata.read_text())["true_encryption"] is True
    assert set(tmp_path.iterdir()) == {source} | ({metadata} if metadata.exists() else set())
